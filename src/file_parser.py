"""File parsing utilities for different file formats."""

import io
import os
import re
from pathlib import Path
from typing import Optional, Tuple
import logging

# Document parsers
try:
    import PyPDF2
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    from docx import Document as DocxDocument
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    from pptx import Presentation
    PPTX_SUPPORT = True
except ImportError:
    PPTX_SUPPORT = False

try:
    import openpyxl
    XLSX_SUPPORT = True
except ImportError:
    XLSX_SUPPORT = False

try:
    import markdown
    from bs4 import BeautifulSoup
    MARKDOWN_SUPPORT = True
except ImportError:
    MARKDOWN_SUPPORT = False

# Text processing
import chardet

from config import settings
from models import FileType

logger = logging.getLogger(__name__)


class FileParser:
    """Parser for various file formats."""
    
    # Maximum file size (100MB)
    MAX_FILE_SIZE = 100 * 1024 * 1024
    
    # Maximum content length
    MAX_CONTENT_LENGTH = 10 * 1024 * 1024  # 10MB of text
    
    @staticmethod
    def get_file_type(extension: str) -> FileType:
        """Determine file type from extension."""
        ext = extension.lower()
        
        document_exts = {'.txt', '.md', '.markdown', '.pdf', '.doc', '.docx', '.ppt', '.pptx', '.xls', '.xlsx'}
        code_exts = {'.py', '.js', '.ts', '.jsx', '.tsx', '.java', '.kt', '.scala', '.c', '.cpp', '.h', '.hpp',
                     '.go', '.rs', '.rb', '.php', '.html', '.htm', '.css', '.scss', '.sass', '.json', '.xml',
                     '.yaml', '.yml', '.sql'}
        config_exts = {'.env', '.ini', '.cfg', '.conf', '.sh', '.bash', '.zsh', '.ps1', '.dockerfile', '.gitignore'}
        
        if ext in document_exts:
            return FileType.DOCUMENT
        elif ext in code_exts:
            return FileType.CODE
        elif ext in config_exts:
            return FileType.CONFIG
        else:
            return FileType.OTHER
    
    @classmethod
    def parse_file(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """
        Parse a file and extract text content.
        
        Returns:
            Tuple of (content, error_message)
        """
        path = Path(filepath)
        
        if not path.exists():
            return None, f"File not found: {filepath}"
        
        if not path.is_file():
            return None, f"Not a file: {filepath}"
        
        file_size = path.stat().st_size
        if file_size > cls.MAX_FILE_SIZE:
            return None, f"File too large: {file_size} bytes (max {cls.MAX_FILE_SIZE})"
        
        if file_size == 0:
            return None, "Empty file"
        
        extension = path.suffix.lower()
        
        try:
            if extension == '.pdf':
                return cls._parse_pdf(filepath)
            elif extension in ['.docx', '.doc']:
                return cls._parse_docx(filepath)
            elif extension in ['.pptx', '.ppt']:
                return cls._parse_pptx(filepath)
            elif extension in ['.xlsx', '.xls']:
                return cls._parse_xlsx(filepath)
            elif extension in ['.md', '.markdown']:
                return cls._parse_markdown(filepath)
            elif extension in ['.html', '.htm']:
                return cls._parse_html(filepath)
            else:
                # Treat as text file
                return cls._parse_text(filepath)
        except Exception as e:
            logger.error(f"Error parsing {filepath}: {e}")
            return None, str(e)
    
    @classmethod
    def _parse_pdf(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse PDF file."""
        if not PDF_SUPPORT:
            return None, "PyPDF2 not installed"
        
        try:
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                
                # Limit content length
                if len(text) > cls.MAX_CONTENT_LENGTH:
                    text = text[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
                
                return text, None
        except Exception as e:
            return None, f"PDF parsing error: {e}"
    
    @classmethod
    def _parse_docx(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse Word document."""
        if not DOCX_SUPPORT:
            return None, "python-docx not installed"
        
        try:
            doc = DocxDocument(filepath)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            
            if len(text) > cls.MAX_CONTENT_LENGTH:
                text = text[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
            
            return text, None
        except Exception as e:
            return None, f"DOCX parsing error: {e}"
    
    @classmethod
    def _parse_pptx(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse PowerPoint presentation."""
        if not PPTX_SUPPORT:
            return None, "python-pptx not installed"
        
        try:
            prs = Presentation(filepath)
            texts = []
            
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
            
            text = "\n".join(texts)
            
            if len(text) > cls.MAX_CONTENT_LENGTH:
                text = text[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
            
            return text, None
        except Exception as e:
            return None, f"PPTX parsing error: {e}"
    
    @classmethod
    def _parse_xlsx(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse Excel spreadsheet."""
        if not XLSX_SUPPORT:
            return None, "openpyxl not installed"
        
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            texts = []
            
            for sheet in wb.worksheets:
                texts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        texts.append(row_text)
            
            text = "\n".join(texts)
            
            if len(text) > cls.MAX_CONTENT_LENGTH:
                text = text[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
            
            return text, None
        except Exception as e:
            return None, f"XLSX parsing error: {e}"
    
    @classmethod
    def _parse_markdown(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse Markdown file."""
        content, error = cls._parse_text(filepath)
        if error:
            return None, error
        
        if MARKDOWN_SUPPORT:
            try:
                # Convert to HTML then extract text
                html = markdown.markdown(content)
                soup = BeautifulSoup(html, 'html.parser')
                text = soup.get_text()
                return text, None
            except:
                # Fallback to raw markdown
                pass
        
        return content, None
    
    @classmethod
    def _parse_html(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse HTML file."""
        content, error = cls._parse_text(filepath)
        if error:
            return None, error
        
        try:
            soup = BeautifulSoup(content, 'html.parser')
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            text = soup.get_text()
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = "\n".join(chunk for chunk in chunks if chunk)
            return text, None
        except:
            return content, None
    
    @classmethod
    def _parse_text(cls, filepath: str) -> Tuple[Optional[str], Optional[str]]:
        """Parse text file with encoding detection."""
        try:
            # Try UTF-8 first
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    content = f.read()
                    if len(content) > cls.MAX_CONTENT_LENGTH:
                        content = content[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
                    return content, None
            except UnicodeDecodeError:
                pass
            
            # Detect encoding
            with open(filepath, 'rb') as f:
                raw_data = f.read()
                result = chardet.detect(raw_data)
                encoding = result.get('encoding', 'utf-8') or 'utf-8'
            
            # Read with detected encoding
            with open(filepath, 'r', encoding=encoding, errors='ignore') as f:
                content = f.read()
                if len(content) > cls.MAX_CONTENT_LENGTH:
                    content = content[:cls.MAX_CONTENT_LENGTH] + "\n[Content truncated...]"
                return content, None
                
        except Exception as e:
            return None, f"Text parsing error: {e}"
    
    @classmethod
    def is_supported(cls, filepath: str) -> bool:
        """Check if file format is supported."""
        path = Path(filepath)
        extension = path.suffix.lower()
        return extension in settings.SUPPORTED_EXTENSIONS
